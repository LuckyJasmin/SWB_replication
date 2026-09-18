"""
05_attention.py
Section 4.3.3 — Perceived similarity and visual attention.

Outputs:
  output/figure4_attention.png   (Figure 4: panel a valence effect; panel b by similarity quartile)
  output/05_tableA6.docx         (Table A6: attention, evaluator fixed-effects estimates)

Run from inside the replication/ folder:  python 05_attention.py
Requires: python-docx
"""
import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib as mpl
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from olstools import ols_cluster

exec(open("00_data_prep.py").read())


def star(p):
    return "***" if p < .001 else "**" if p < .01 else "*" if p < .05 else ""


def demean(v, g):
    v = np.asarray(v, float)
    return v - pd.Series(v).groupby(g).transform("mean").values


df, subj = build_data()
sh = judgment_panel(df, subj)

ctx = ["age", "gender", "cld", "fml", "income", "health"]
sh["fix_context"] = sh[[f"fix_ttldrt_{e}" for e in ctx]].sum(axis=1)
sh["fix_shock"] = sh["fix_ttldrt_sk"]
sh["fix_total"] = sh["fix_context"] + sh["fix_shock"] + sh["fix_ttldrt_qn"]
for v in ["fix_total", "fix_context", "fix_shock"]:
    sh[f"l_{v}"] = np.log1p(sh[v])


def fe(dv):
    d = sh.dropna(subset=[dv, "neg", "perc_z"]).copy()
    g = d["SubjectID"].values
    d["int"] = d["neg"] * d["perc_z"]
    y = demean(d[dv], g)
    X = np.column_stack([demean(d["neg"], g), demean(d["int"], g)])
    r, m = ols_cluster(y, X, g, ["Negative", "Negative x PercSim"])
    m["subj"] = d["SubjectID"].nunique()
    return r, m


RES = {lab: fe(dv) for dv, lab in
       [("l_fix_total", "Total"), ("l_fix_context", "Contextual profile"), ("l_fix_shock", "Shock text")]}


def neg_effect(dv, group):
    d = sh[sh["grp"] == group].dropna(subset=[dv, "neg"]).copy()
    g = d["SubjectID"].values
    y = demean(d[dv], g); X = demean(d["neg"], g).reshape(-1, 1)
    r, m = ols_cluster(y, X, g, ["Negative"])
    return r.loc["Negative", "coef"], r.loc["Negative", "se"]


# ============================================================
# FIGURE 4  (unchanged)
# ============================================================
q25, q75 = subj["s_smlrt"].quantile(.25), subj["s_smlrt"].quantile(.75)
sh["grp"] = np.where(sh["s_smlrt"] <= q25, "dissimilar",
                     np.where(sh["s_smlrt"] >= q75, "similar", None))

mpl.rcParams.update({"font.family": "serif", "font.serif": ["DejaVu Serif"],
    "axes.spines.top": False, "axes.spines.right": False,
    "axes.grid": True, "grid.alpha": 0.22, "savefig.dpi": 200, "savefig.bbox": "tight"})
A_CTX, A_SHK, A_TOT = "#4E7D6E", "#B8853A", "#7A6A9E"
DIS, SIM = "#A8C4B8", "#3B6B5E"

fig, axes = plt.subplots(1, 2, figsize=(12.8, 4.8))
ax = axes[0]
cats = ["Contextual\nprofile", "Shock\ntext", "Total\nfixation"]
vals = [RES["Contextual profile"][0].loc["Negative", "coef"],
        RES["Shock text"][0].loc["Negative", "coef"],
        RES["Total"][0].loc["Negative", "coef"]]
ses = [RES["Contextual profile"][0].loc["Negative", "se"],
       RES["Shock text"][0].loc["Negative", "se"],
       RES["Total"][0].loc["Negative", "se"]]
cols = [A_CTX, A_SHK, A_TOT]
x = np.array([0, 1.3, 2.6])
ax.axhline(0, color="k", lw=1)
for i in range(3):
    ax.errorbar(x[i], vals[i], yerr=1.96*ses[i], fmt="o", color=cols[i], ms=9, elinewidth=2, capsize=4)
ax.set_xticks(x); ax.set_xticklabels(cats); ax.set_xlim(-0.7, 3.3)
ax.set_ylabel("Effect of a negative vs. positive\nshock on fixation (log ms)")
ax.set_title("(a) Negative shocks shift attention to the event", weight="bold", pad=10)

ax = axes[1]
groups = ["dissimilar", "similar"]
glab = ["Most dissimilar\n(bottom quartile)", "Most similar\n(top quartile)"]
b = [neg_effect("l_fix_shock", gg)[0] for gg in groups]
se = [neg_effect("l_fix_shock", gg)[1] for gg in groups]
xb = np.arange(2)
ax.axhline(0, color="k", lw=1)
for i, cc in enumerate([DIS, SIM]):
    ax.bar(xb[i], b[i], width=0.55, color=cc, edgecolor="white",
           yerr=1.96*se[i], error_kw=dict(ecolor="#444", lw=1.3, capsize=4))
ax.set_xticks(xb); ax.set_xticklabels(glab)
ax.set_ylabel("Effect of a negative vs. positive shock\non shock-text fixation (log ms)")
ax.set_title("(b) …more so for dissimilar targets", weight="bold", pad=10)
fig.subplots_adjust(wspace=0.34, bottom=0.16)
os.makedirs("output", exist_ok=True)
fig.savefig("output/figure4_attention.png"); plt.close(fig)
print("saved output/figure4_attention.png")

# ============================================================
# TABLE A6 — attention, evaluator fixed-effects estimates
# ============================================================
doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(11)


def setc(cell, txt, bold=False, ital=False, align="center", size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    r = p.add_run(str(txt)); r.bold = bold; r.italic = ital
    r.font.name = "Times New Roman"; r.font.size = Pt(size)


p = doc.add_paragraph()
rr = p.add_run("Table A6. Attention and perceived similarity, evaluator fixed-effects estimates")
rr.bold = True; rr.font.size = Pt(11); rr.font.name = "Times New Roman"

cols_order = ["Total", "Contextual profile", "Shock text"]
t = doc.add_table(rows=1, cols=4); t.alignment = WD_TABLE_ALIGNMENT.CENTER
h = t.rows[0].cells
setc(h[0], "", bold=True, align="left")
for j, lab in enumerate(cols_order):
    setc(h[j+1], lab, bold=True)


def row(label, name):
    cr = t.add_row().cells; setc(cr[0], label, align="left")
    sr = t.add_row().cells; setc(sr[0], "", align="left")
    for j, lab in enumerate(cols_order):
        r, m = RES[lab]
        setc(cr[j+1], f"{r.loc[name,'coef']:.3f}{star(r.loc[name,'p'])}")
        setc(sr[j+1], f"({r.loc[name,'se']:.3f})", size=8)


row("Negative shock", "Negative")
row("Negative × Perceived similarity", "Negative x PercSim")

rr = t.add_row().cells; setc(rr[0], "Evaluator fixed effects", align="left", size=8)
for j in range(3):
    setc(rr[j+1], "Yes", size=8)
rr = t.add_row().cells; setc(rr[0], "Observations", align="left", size=8)
for j, lab in enumerate(cols_order):
    setc(rr[j+1], f"{int(RES[lab][1]['N']):,}", size=8)
rr = t.add_row().cells; setc(rr[0], "Evaluators", align="left", size=8)
for j, lab in enumerate(cols_order):
    setc(rr[j+1], f"{int(RES[lab][1]['subj'])}", size=8)

p = doc.add_paragraph(); nr = p.add_run(
 "Notes: Evaluator fixed-effects estimates on the judgment-task eye-tracking panel. The dependent "
 "variable in Column (1) is total fixation duration across all areas of interest; Column (2) uses "
 "fixation duration on the target’s contextual profile; and Column (3) uses fixation duration on the "
 "focal shock text. Each fixation outcome is measured in log milliseconds (natural log of one plus "
 "total fixation duration). Perceived similarity is standardized and constant within evaluator. "
 "Cluster-robust standard errors (by evaluator) in parentheses. * p<0.05, ** p<0.01, *** p<0.001.")
nr.italic = True; nr.font.size = Pt(8); nr.font.name = "Times New Roman"

doc.save("output/05_tableA6.docx")
print("saved output/05_tableA6.docx")