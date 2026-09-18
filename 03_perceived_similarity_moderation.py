"""
03_perceived_similarity_moderation.py
Writes Table 2 (Perceived similarity and the negativity bias) to a Word file:
  output/03_table2.docx

Columns:
  (1) ΔSWB ~ Negative                          [loss-aversion baseline]
  (2) + raw 0-10 perceived similarity + interaction
  (3) + standardized similarity + interaction
  (4) (3) + evaluator controls
  (5) (3) + evaluator fixed effects

Run from inside the replication/ folder:  python 03_perceived_similarity_moderation.py
Requires: python-docx
"""
import os
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from olstools import ols_cluster

exec(open("00_data_prep.py").read())


def star(p):
    return "***" if p < .001 else "**" if p < .01 else "*" if p < .05 else ""


def demean(v, g):
    v = np.asarray(v, float)
    return v - pd.Series(v).groupby(g).transform("mean").values


# ============================================================
# 1. ESTIMATION
# ============================================================
df, subj = build_data()
sh = judgment_panel(df, subj)
COVS = ["ev_female", "ev_married", "ev_loginc", "inc_miss",
        "ev_age", "ev_edu", "ev_nonwhite", "ev_ownswb"]

d = sh.dropna(subset=["dswb", "neg", "perc", "perc0", "perc_z"]).copy()
g = d["SubjectID"].values

# Col 1: Negative only
X = np.column_stack([np.ones(len(d)), d["neg"]])
r1, m1 = ols_cluster(d["dswb"].values, X, g, ["const", "Negative"])

# Col 2: raw 0-10 similarity + interaction
d["int0"] = d["neg"] * d["perc0"]
X = np.column_stack([np.ones(len(d)), d["neg"], d["perc0"], d["int0"]])
r2, m2 = ols_cluster(d["dswb"].values, X, g,
                     ["const", "Negative", "PercSim", "NegxPercSim"])

# Col 3: standardized similarity + interaction
d["intz"] = d["neg"] * d["perc_z"]
X = np.column_stack([np.ones(len(d)), d["neg"], d["perc_z"], d["intz"]])
r3, m3 = ols_cluster(d["dswb"].values, X, g,
                     ["const", "Negative", "PercSim", "NegxPercSim"])

# Col 4: + evaluator controls
dc = d.dropna(subset=COVS).copy()
gc = dc["SubjectID"].values
dc["intz"] = dc["neg"] * dc["perc_z"]
X = np.column_stack([np.ones(len(dc)), dc["neg"], dc["perc_z"], dc["intz"]] +
                    [dc[c].values for c in COVS])
r4, m4 = ols_cluster(dc["dswb"].values, X, gc,
                     ["const", "Negative", "PercSim", "NegxPercSim"] + COVS)

# Col 5: + evaluator fixed effects (within-participant demeaning)
d5 = d.copy()
d5["intz"] = d5["neg"] * d5["perc_z"]
y = demean(d5["dswb"], g)
X = np.column_stack([demean(d5["neg"], g), demean(d5["intz"], g)])
r5, m5 = ols_cluster(y, X, g, ["Negative", "NegxPercSim"])
# within R2 (matches the manuscript's reported R2 for the FE column reasonably;
# reported as the model R2 from the demeaned fit)
r5_r2 = m5["r2"]

COLS = [(r1, m1), (r2, m2), (r3, m3), (r4, m4), (r5, m5)]

# ============================================================
# 2. WORD OUTPUT
# ============================================================
doc = Document()
doc.styles["Normal"].font.name = "Times New Roman"
doc.styles["Normal"].font.size = Pt(11)


def setc(cell, txt, bold=False, ital=False, align="center", size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    r = p.add_run(str(txt)); r.bold = bold; r.italic = ital
    r.font.name = "Times New Roman"; r.font.size = Pt(size)


p = doc.add_paragraph(); rr = p.add_run("Table 2. Perceived similarity and the negativity bias")
rr.bold = True; rr.font.size = Pt(11); rr.font.name = "Times New Roman"

t = doc.add_table(rows=1, cols=6); t.alignment = WD_TABLE_ALIGNMENT.CENTER
h = t.rows[0].cells
setc(h[0], "Dependent variable: ΔSWB", bold=True, align="left", size=8)
for j in range(5):
    setc(h[j+1], f"({j+1})", bold=True)


def coef_row(label, getter):
    cr = t.add_row().cells; setc(cr[0], label, align="left")
    sr = t.add_row().cells; setc(sr[0], "", align="left")
    for j, res in enumerate(COLS):
        c, s = getter(res, j)
        setc(cr[j+1], c); setc(sr[j+1], s, size=8)


def get_neg(res, j):
    r, m = res
    if "Negative" in r.index:
        return f"{r.loc['Negative','coef']:.3f}{star(r.loc['Negative','p'])}", f"({r.loc['Negative','se']:.3f})"
    return "", ""


def get_perc(res, j):
    r, m = res
    # only columns 2,3,4 have a similarity main effect (col1 none, col5 absorbed by FE)
    if "PercSim" in r.index:
        return f"{r.loc['PercSim','coef']:.3f}{star(r.loc['PercSim','p'])}", f"({r.loc['PercSim','se']:.3f})"
    return "", ""


def get_int(res, j):
    r, m = res
    if "NegxPercSim" in r.index:
        return f"{r.loc['NegxPercSim','coef']:.3f}{star(r.loc['NegxPercSim','p'])}", f"({r.loc['NegxPercSim','se']:.3f})"
    return "", ""


coef_row("Negative shock", get_neg)
coef_row("Perceived similarity", get_perc)
coef_row("Negative × Perceived similarity", get_int)

# meta rows
rr = t.add_row().cells; setc(rr[0], "Evaluator controls", align="left", size=8)
for j, flag in enumerate(["No", "No", "No", "Yes", "No"]):
    setc(rr[j+1], flag, size=8)
rr = t.add_row().cells; setc(rr[0], "Evaluator fixed effects", align="left", size=8)
for j, flag in enumerate(["No", "No", "No", "No", "Yes"]):
    setc(rr[j+1], flag, size=8)
rr = t.add_row().cells; setc(rr[0], "Observations", align="left", size=8)
for j, res in enumerate(COLS):
    setc(rr[j+1], f"{int(res[1]['N']):,}", size=8)
rr = t.add_row().cells; setc(rr[0], "R²", align="left", size=8)
for j, res in enumerate(COLS):
    setc(rr[j+1], f"{res[1]['r2']:.3f}", size=8)

p = doc.add_paragraph(); nr = p.add_run(
 "Notes: Dependent variable is the change in the target’s judged life satisfaction (ΔSWB). The unit "
 "of observation is the evaluator × shock (8 shocks per evaluator). Cluster-robust standard errors "
 "(clustered by participant) in parentheses. Column 2 uses the raw 0–10 similarity scale; columns "
 "3–5 standardize similarity to mean 0, SD 1, so the interaction is the change in the negativity "
 "bias per standard deviation of perceived similarity. Column 4 adds evaluator controls (gender, "
 "marital status, log individual income and a missing-income indicator, age, education, a non-white "
 "indicator, and own life satisfaction). Column 5 includes evaluator fixed effects. * p<0.05, "
 "** p<0.01, *** p<0.001.")
nr.italic = True; nr.font.size = Pt(8); nr.font.name = "Times New Roman"

os.makedirs("output", exist_ok=True)
doc.save("output/03_table2.docx")
print("saved output/03_table2.docx")